"""Motor de reportes — CU-44 (admin_tenant) y CU-45 (cliente) — con 3 formatos.

Diseño: las funciones `datos_reporte_*` construyen un "documento" genérico
(dict con título, metadatos y secciones) y los serializadores `to_html` / `to_pdf`
/ `to_excel` lo convierten al formato pedido. Así web y móvil reutilizan el mismo
motor y un solo endpoint puede emitir PDF, Excel o HTML.

Estructura del documento:
{
  "titulo":    str,
  "subtitulo": str | None,
  "meta":      [(etiqueta, valor), ...],          # datos clave arriba
  "secciones": [
     {"titulo": str, "tipo": "tabla", "columnas": [...], "filas": [[...], ...]},
     {"titulo": str, "tipo": "kv",    "items": [(etiqueta, valor), ...]},
  ],
}
"""
from __future__ import annotations

import io
from datetime import date, datetime
from html import escape
from uuid import UUID

from sqlalchemy.orm import Session

from app.core.tenant_context import aplicar_filtro_tenant
from app.core.timezone import now_bo
from app.models.asignacion import Asignacion
from app.models.calificacion import Calificacion
from app.models.incidente import Incidente
from app.models.pago import Pago
from app.models.taller import Taller
from app.models.usuario import Usuario
from app.models.vehiculo import Vehiculo
from app.services import kpi_service

FORMATOS_VALIDOS = ("pdf", "excel", "html")

_ESTADO_LABEL = {
    "pendiente": "Pendiente",
    "buscando_taller": "Buscando taller",
    "taller_asignado": "Taller asignado",
    "en_camino": "En camino",
    "en_atencion": "En atención",
    "finalizado": "Finalizado",
    "cancelado": "Cancelado",
}


def _fmt_fecha(dt: datetime | None) -> str:
    if dt is None:
        return "—"
    return dt.strftime("%d/%m/%Y %H:%M")


def _rango_label(desde: date | None, hasta: date | None) -> str:
    if desde and hasta:
        return f"{desde.strftime('%d/%m/%Y')} – {hasta.strftime('%d/%m/%Y')}"
    if desde:
        return f"Desde {desde.strftime('%d/%m/%Y')}"
    if hasta:
        return f"Hasta {hasta.strftime('%d/%m/%Y')}"
    return "Todo el histórico"


# ═══════════════════════════════════════════════════════════════════════════
# CU-45 — Reporte de servicios del cliente
# ═══════════════════════════════════════════════════════════════════════════

def datos_reporte_cliente(
    db: Session,
    cliente: Usuario,
    desde: date | None = None,
    hasta: date | None = None,
    estado: str | None = None,
    incluir_resumen: bool = True,
    incluir_detalle: bool = True,
) -> dict:
    """Construye el documento del reporte de servicios del cliente.

    `estado`: filtro opcional ('finalizado' / 'cancelado'); None = todos.
    `incluir_resumen` / `incluir_detalle`: el cliente personaliza qué secciones ve.
    """
    q = db.query(Incidente).filter(Incidente.cliente_id == cliente.id)
    if desde:
        q = q.filter(Incidente.creado_en >= datetime.combine(desde, datetime.min.time()))
    if hasta:
        q = q.filter(Incidente.creado_en <= datetime.combine(hasta, datetime.max.time()))
    if estado:
        q = q.filter(Incidente.estado == estado)
    incidentes = q.order_by(Incidente.creado_en.desc()).all()

    ids = [i.id for i in incidentes]

    # Datos relacionados en lote
    asignaciones: dict[UUID, Asignacion] = {}
    pagos: dict[UUID, Pago] = {}
    calificaciones: dict[UUID, Calificacion] = {}
    vehiculos: dict[UUID, Vehiculo] = {}
    talleres: dict[UUID, Taller] = {}
    if ids:
        for a in (
            db.query(Asignacion)
            .filter(Asignacion.incidente_id.in_(ids), Asignacion.accion_taller == "aceptado")
            .all()
        ):
            asignaciones[a.incidente_id] = a
        for p in db.query(Pago).filter(Pago.incidente_id.in_(ids)).all():
            pagos[p.incidente_id] = p
        for c in db.query(Calificacion).filter(Calificacion.incidente_id.in_(ids)).all():
            calificaciones[c.incidente_id] = c
        veh_ids = [i.vehiculo_id for i in incidentes if i.vehiculo_id]
        if veh_ids:
            for v in db.query(Vehiculo).filter(Vehiculo.id.in_(veh_ids)).all():
                vehiculos[v.id] = v
        taller_ids = [a.taller_id for a in asignaciones.values()]
        if taller_ids:
            for t in db.query(Taller).filter(Taller.id.in_(taller_ids)).all():
                talleres[t.id] = t

    # Filas de detalle
    filas: list[list[str]] = []
    total_pagado = 0.0
    suma_calif = 0
    n_calif = 0
    finalizados = 0
    cancelados = 0
    for inc in incidentes:
        if inc.estado == "finalizado":
            finalizados += 1
        elif inc.estado == "cancelado":
            cancelados += 1

        veh = vehiculos.get(inc.vehiculo_id) if inc.vehiculo_id else None
        veh_txt = f"{veh.marca} {veh.modelo} ({veh.placa})" if veh else "—"

        asig = asignaciones.get(inc.id)
        taller = talleres.get(asig.taller_id) if asig else None
        taller_txt = taller.nombre if taller else "—"

        pago = pagos.get(inc.id)
        if pago and pago.estado == "pagado":
            total_pagado += float(pago.monto_total)
        monto_txt = f"Bs. {float(pago.monto_total):.2f}" if pago else "—"

        cal = calificaciones.get(inc.id)
        if cal:
            suma_calif += cal.puntuacion
            n_calif += 1
        cal_txt = f"{cal.puntuacion}★" if cal else "—"

        filas.append([
            _fmt_fecha(inc.creado_en),
            veh_txt,
            (inc.clasificacion_ia or "—").capitalize(),
            _ESTADO_LABEL.get(inc.estado, inc.estado),
            taller_txt,
            monto_txt,
            cal_txt,
        ])

    secciones: list[dict] = []

    if incluir_resumen:
        prom_txt = f"{round(suma_calif / n_calif, 2)}★" if n_calif else "—"
        secciones.append({
            "titulo": "Resumen",
            "tipo": "kv",
            "items": [
                ("Total de emergencias", str(len(incidentes))),
                ("Servicios finalizados", str(finalizados)),
                ("Servicios cancelados", str(cancelados)),
                ("Total pagado", f"Bs. {total_pagado:.2f}"),
                ("Calificación promedio que diste", prom_txt),
            ],
        })

    if incluir_detalle:
        secciones.append({
            "titulo": "Detalle de servicios",
            "tipo": "tabla",
            "columnas": ["Fecha", "Vehículo", "Tipo", "Estado", "Taller", "Monto", "Calif."],
            "filas": filas,
        })

    return {
        "titulo": "Reporte de mis servicios",
        "subtitulo": cliente.nombre_completo,
        "meta": [
            ("Período", _rango_label(desde, hasta)),
            ("Filtro de estado", _ESTADO_LABEL.get(estado, "Todos") if estado else "Todos"),
            ("Generado", _fmt_fecha(now_bo().replace(tzinfo=None))),
        ],
        "secciones": secciones,
    }


# ═══════════════════════════════════════════════════════════════════════════
# CU-44 — Reporte operacional del tenant (admin_tenant)
# ═══════════════════════════════════════════════════════════════════════════

def datos_reporte_admin(
    db: Session,
    current_user: Usuario,
    desde: date | None = None,
    hasta: date | None = None,
    secciones_sel: set[str] | None = None,
) -> dict:
    """Reporte operacional del tenant, reusando el cálculo de KPIs (CU-39).

    `secciones_sel`: subconjunto de
    {'resumen','por_tipo','talleres','zonas','sla','satisfaccion'} — personalización.
    None = todas.
    """
    todas = {"resumen", "por_tipo", "talleres", "zonas", "sla", "satisfaccion"}
    sel = secciones_sel if secciones_sel else todas

    k = kpi_service.calcular_kpis(db, current_user, desde, hasta)
    secciones: list[dict] = []

    def _min(v) -> str:
        return f"{v} min" if v is not None else "—"

    def _pct(v) -> str:
        return f"{v}%" if v is not None else "—"

    if "resumen" in sel:
        secciones.append({
            "titulo": "Resumen operacional",
            "tipo": "kv",
            "items": [
                ("Incidentes en el período", str(k["totales"]["incidentes"])),
                ("Con taller asignado", str(k["totales"]["cotizaciones_aceptadas"])),
                ("Tiempo promedio de asignación", _min(k["tiempo_promedio_asignacion_min"])),
                ("Tiempo promedio de llegada", _min(k["tiempo_promedio_llegada_min"])),
                ("Casos cancelados", f'{k["casos_cancelados"]["cantidad"]} ({k["casos_cancelados"]["tasa_porcentaje"]}%)'),
                ("Cumplimiento de SLA", _pct(k["cumplimiento_sla"]["porcentaje"])),
                ("Satisfacción promedio", f'{k["satisfaccion"]["promedio"]}★ ({k["satisfaccion"]["total_resenas"]} reseñas)'
                    if k["satisfaccion"]["promedio"] is not None else "—"),
            ],
        })

    if "por_tipo" in sel:
        secciones.append({
            "titulo": "Incidentes por tipo",
            "tipo": "tabla",
            "columnas": ["Tipo", "Cantidad"],
            "filas": [[str(t).capitalize(), str(c)] for t, c in k["incidentes_por_tipo"].items()],
        })

    if "talleres" in sel:
        secciones.append({
            "titulo": "Talleres más eficientes",
            "tipo": "tabla",
            "columnas": ["Taller", "Completados", "Tiempo prom. (min)"],
            "filas": [
                [t["nombre"], str(t["incidentes_completados"]), str(t["tiempo_promedio_min"])]
                for t in k["talleres_mas_eficientes"]
            ],
        })

    if "zonas" in sel:
        secciones.append({
            "titulo": "Zonas con más incidentes",
            "tipo": "tabla",
            "columnas": ["Latitud", "Longitud", "Incidentes"],
            "filas": [
                [str(z["latitud"]), str(z["longitud"]), str(z["incidentes"])]
                for z in k["zonas_con_mas_incidentes"]
            ],
        })

    if "sla" in sel:
        secciones.append({
            "titulo": "Cumplimiento de SLA por tipo",
            "tipo": "tabla",
            "columnas": ["Tipo", "Aplicables", "Cumplidos", "%"],
            "filas": [
                [str(tipo).capitalize(), str(v["aplicables"]), str(v["cumplidos"]), _pct(v["porcentaje"])]
                for tipo, v in k["cumplimiento_sla"]["por_tipo"].items()
            ],
        })

    if "satisfaccion" in sel:
        s = k["satisfaccion"]
        secciones.append({
            "titulo": "Satisfacción de clientes",
            "tipo": "kv",
            "items": [
                ("Promedio", f'{s["promedio"]}★' if s["promedio"] is not None else "—"),
                ("Total de reseñas", str(s["total_resenas"])),
            ],
        })

    return {
        "titulo": "Reporte operacional",
        "subtitulo": "Red de talleres",
        "meta": [
            ("Período", _rango_label(desde, hasta)),
            ("Generado", _fmt_fecha(now_bo().replace(tzinfo=None))),
        ],
        "secciones": secciones,
    }


# ═══════════════════════════════════════════════════════════════════════════
# Serializadores
# ═══════════════════════════════════════════════════════════════════════════

def to_html(doc: dict) -> bytes:
    p: list[str] = []
    p.append("<!DOCTYPE html><html lang='es'><head><meta charset='utf-8'>")
    p.append(f"<title>{escape(doc['titulo'])}</title>")
    p.append(
        "<style>"
        "body{font-family:Arial,Helvetica,sans-serif;color:#1a1a1a;margin:32px;}"
        "h1{color:#E63946;margin-bottom:2px;}h2{margin-top:28px;border-bottom:2px solid #eee;padding-bottom:4px;}"
        ".sub{color:#666;margin-top:0;}"
        ".meta{margin:16px 0;font-size:14px;color:#444;}"
        ".meta span{display:inline-block;margin-right:18px;}"
        "table{border-collapse:collapse;width:100%;margin-top:8px;font-size:13px;}"
        "th,td{border:1px solid #ddd;padding:7px 10px;text-align:left;}"
        "th{background:#f4f4f6;}tr:nth-child(even){background:#fafafa;}"
        ".kv{font-size:14px;}.kv div{padding:4px 0;}"
        ".kv b{display:inline-block;min-width:260px;color:#555;}"
        "</style></head><body>"
    )
    p.append(f"<h1>{escape(doc['titulo'])}</h1>")
    if doc.get("subtitulo"):
        p.append(f"<p class='sub'>{escape(doc['subtitulo'])}</p>")
    if doc.get("meta"):
        p.append("<div class='meta'>")
        for etq, val in doc["meta"]:
            p.append(f"<span><b>{escape(str(etq))}:</b> {escape(str(val))}</span>")
        p.append("</div>")

    for sec in doc["secciones"]:
        p.append(f"<h2>{escape(sec['titulo'])}</h2>")
        if sec["tipo"] == "kv":
            p.append("<div class='kv'>")
            for etq, val in sec["items"]:
                p.append(f"<div><b>{escape(str(etq))}</b> {escape(str(val))}</div>")
            p.append("</div>")
        else:  # tabla
            p.append("<table><thead><tr>")
            for col in sec["columnas"]:
                p.append(f"<th>{escape(str(col))}</th>")
            p.append("</tr></thead><tbody>")
            if not sec["filas"]:
                p.append(f"<tr><td colspan='{len(sec['columnas'])}' style='color:#999'>Sin datos</td></tr>")
            for fila in sec["filas"]:
                p.append("<tr>" + "".join(f"<td>{escape(str(c))}</td>" for c in fila) + "</tr>")
            p.append("</tbody></table>")

    p.append("</body></html>")
    return "".join(p).encode("utf-8")


def to_pdf(doc: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib.units import cm
    from reportlab.platypus import (
        SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle,
    )

    buf = io.BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4,
                            topMargin=1.6 * cm, bottomMargin=1.6 * cm,
                            leftMargin=1.6 * cm, rightMargin=1.6 * cm)
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("h1c", parent=styles["Title"], textColor=colors.HexColor("#E63946"))
    sub = ParagraphStyle("subc", parent=styles["Normal"], textColor=colors.grey, fontSize=10)
    h2 = ParagraphStyle("h2c", parent=styles["Heading2"], spaceBefore=14)

    elems: list = [Paragraph(escape(doc["titulo"]), h1)]
    if doc.get("subtitulo"):
        elems.append(Paragraph(escape(doc["subtitulo"]), sub))
    if doc.get("meta"):
        meta_txt = "  •  ".join(f"<b>{escape(str(e))}:</b> {escape(str(v))}" for e, v in doc["meta"])
        elems.append(Spacer(1, 6))
        elems.append(Paragraph(meta_txt, sub))
    elems.append(Spacer(1, 10))

    for sec in doc["secciones"]:
        elems.append(Paragraph(escape(sec["titulo"]), h2))
        if sec["tipo"] == "kv":
            data = [[str(e), str(v)] for e, v in sec["items"]]
            if not data:
                data = [["Sin datos", ""]]
            t = Table(data, colWidths=[8 * cm, 8 * cm])
            t.setStyle(TableStyle([
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("TEXTCOLOR", (0, 0), (0, -1), colors.HexColor("#555555")),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ("LINEBELOW", (0, 0), (-1, -1), 0.3, colors.HexColor("#eeeeee")),
            ]))
            elems.append(t)
        else:
            data = [list(map(str, sec["columnas"]))]
            data += [[str(c) for c in fila] for fila in sec["filas"]]
            if len(data) == 1:
                data.append(["Sin datos"] + [""] * (len(sec["columnas"]) - 1))
            t = Table(data, repeatRows=1)
            t.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f4f4f6")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.HexColor("#333333")),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#dddddd")),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#fafafa")]),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ]))
            elems.append(t)

    pdf.build(elems)
    return buf.getvalue()


def to_excel(doc: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Reporte"

    bold = Font(bold=True)
    titulo_font = Font(bold=True, size=14, color="E63946")
    head_fill = PatternFill("solid", fgColor="F4F4F6")

    fila = 1
    ws.cell(row=fila, column=1, value=doc["titulo"]).font = titulo_font
    fila += 1
    if doc.get("subtitulo"):
        ws.cell(row=fila, column=1, value=doc["subtitulo"])
        fila += 1
    for etq, val in doc.get("meta", []):
        ws.cell(row=fila, column=1, value=str(etq)).font = bold
        ws.cell(row=fila, column=2, value=str(val))
        fila += 1
    fila += 1

    for sec in doc["secciones"]:
        ws.cell(row=fila, column=1, value=sec["titulo"]).font = Font(bold=True, size=12)
        fila += 1
        if sec["tipo"] == "kv":
            for etq, val in sec["items"]:
                ws.cell(row=fila, column=1, value=str(etq)).font = bold
                ws.cell(row=fila, column=2, value=str(val))
                fila += 1
        else:
            for j, col in enumerate(sec["columnas"], start=1):
                c = ws.cell(row=fila, column=j, value=str(col))
                c.font = bold
                c.fill = head_fill
            fila += 1
            for f in sec["filas"]:
                for j, val in enumerate(f, start=1):
                    ws.cell(row=fila, column=j, value=str(val))
                fila += 1
        fila += 1

    # Ancho de columnas aproximado
    for col_cells in ws.columns:
        largo = max((len(str(c.value)) for c in col_cells if c.value), default=10)
        ws.column_dimensions[col_cells[0].column_letter].width = min(largo + 4, 50)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ═══════════════════════════════════════════════════════════════════════════
# Despacho por formato
# ═══════════════════════════════════════════════════════════════════════════

def generar(doc: dict, formato: str, nombre_base: str) -> tuple[bytes, str, str]:
    """Devuelve (contenido, media_type, filename) según el formato pedido."""
    formato = (formato or "pdf").lower()
    if formato == "html":
        return to_html(doc), "text/html; charset=utf-8", f"{nombre_base}.html"
    if formato == "excel":
        return (
            to_excel(doc),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            f"{nombre_base}.xlsx",
        )
    # default pdf
    return to_pdf(doc), "application/pdf", f"{nombre_base}.pdf"
